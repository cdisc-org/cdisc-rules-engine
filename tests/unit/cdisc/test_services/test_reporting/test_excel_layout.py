import os
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import MagicMock

import openpyxl
import pytest

from cdisc_rules_engine.services.reporting.excel_report import ExcelReport, _format_row_range
from cdisc_rules_engine.services.reporting.excel_writer import (
    excel_open_workbook,
    excel_update_worksheet,
    excel_workbook_to_stream,
    group_rows_by_key,
)

test_report_template: str = f"{os.path.dirname(__file__)}/../../../../../resources/templates/report-template.xlsx"


def open_template():
    with open(test_report_template, "rb") as f:
        return excel_open_workbook(f.read())


def detail_row(core_id, dataset, row, message="a message", values="null"):
    # CORE-ID, Message, Executability, Dataset, USUBJID, Record, Sequence, Variable(s), Value(s)
    return [core_id, message, "fully executable", dataset, f"SUBJ-{row}", row, row, "AESEQ", values]


def report_for(datasets, args=None):
    report = ExcelReport.__new__(ExcelReport)
    report._datasets = datasets
    report._args = args or MagicMock()
    return report


def test_every_written_row_keeps_wrapped_text():
    wb = open_template()
    rows = [detail_row("CORE-000001", "ae.xpt", index) for index in range(1, 400)]

    excel_update_worksheet(wb["Issue Details"], rows, dict(wrap_text=True))

    saved = openpyxl.load_workbook(BytesIO(excel_workbook_to_stream(wb)))["Issue Details"]
    assert all(saved.cell(row=row, column=2).alignment.wrap_text for row in (2, 100, 300, 400))


def test_empty_code_rows_are_still_filled():
    wb = openpyxl.Workbook()
    ws = wb.active

    excel_update_worksheet(ws, [["C1", "", "x"], ["C2", "codelist", "y"]], dict(wrap_text=True), fill_empty_rows=True)

    assert ws.cell(row=2, column=1).fill.fgColor.rgb.lower().endswith("ccffff")
    assert ws.cell(row=3, column=1).fill.fill_type is None
    assert ws.cell(row=2, column=3).alignment.wrap_text


def test_rows_sharing_a_core_id_collapse_under_their_first_row():
    ws = openpyxl.Workbook().active
    rows = [
        detail_row("CORE-1", "ae.xpt", 1),
        detail_row("CORE-1", "ae.xpt", 2),
        detail_row("CORE-1", "ae.xpt", 3),
        detail_row("CORE-2", "ae.xpt", 1),
        detail_row("CORE-3", "ae.xpt", 1),
        detail_row("CORE-3", "ae.xpt", 2),
    ]

    group_rows_by_key(ws, rows, key_col=0)

    levels = {row: ws.row_dimensions[row].outlineLevel for row in range(2, 8)}
    assert levels == {2: 0, 3: 1, 4: 1, 5: 0, 6: 0, 7: 1}
    assert ws.sheet_properties.outlinePr.summaryBelow is False


@pytest.mark.parametrize(
    "values, expected",
    [
        ([1, 2, 3, 5, 7, 8], "1-3, 5, 7-8"),
        (["3", "1", "2", 2], "1-3"),
        ([4, "", None, "A1"], "4, A1"),
        ([], ""),
    ],
)
def test_format_row_range(values, expected):
    assert _format_row_range(values) == expected


def test_an_issue_on_every_record_collapses_to_one_row():
    report = report_for([SimpleNamespace(filename="ae.xpt", record_count=3)])
    rows = [detail_row("CORE-1", "ae.xpt", row) for row in (1, 2, 3)]

    collapsed = report._collapse_duplicate_detail_rows(rows)

    assert len(collapsed) == 1
    assert collapsed[0][4] == ""
    assert collapsed[0][5] == "1-3"
    assert collapsed[0][6] == "1-3"


def test_an_issue_on_some_records_keeps_one_row_per_record():
    report = report_for([SimpleNamespace(filename="ae.xpt", record_count=4)])
    rows = [detail_row("CORE-1", "ae.xpt", row) for row in (1, 2, 3)]

    assert report._collapse_duplicate_detail_rows(rows) == rows


def test_interleaved_issues_keep_their_order_when_nothing_collapses():
    report = report_for([SimpleNamespace(filename="ae.xpt", record_count=4)])
    rows = [
        detail_row("CORE-1", "ae.xpt", 1, values="X"),
        detail_row("CORE-1", "ae.xpt", 2, values="Y"),
        detail_row("CORE-1", "ae.xpt", 3, values="X"),
    ]

    assert report._collapse_duplicate_detail_rows(rows) == rows


def test_a_collapsed_issue_takes_the_place_of_its_first_row():
    report = report_for([SimpleNamespace(filename="ae.xpt", record_count=2)])
    rows = [
        detail_row("CORE-1", "ae.xpt", 1, message="other", values="Y"),
        detail_row("CORE-1", "ae.xpt", 1),
        detail_row("CORE-1", "ae.xpt", 2, message="third", values="Z"),
        detail_row("CORE-1", "ae.xpt", 2),
    ]

    collapsed = report._collapse_duplicate_detail_rows(rows)

    assert [(row[1], row[5]) for row in collapsed] == [("other", 1), ("a message", "1-2"), ("third", 2)]


def test_collapsing_is_off_by_default():
    report = report_for([SimpleNamespace(filename="ae.xpt", record_count=2)], args=SimpleNamespace())
    report.get_detailed_data = lambda excel: [detail_row("CORE-1", "ae.xpt", row) for row in (1, 2)]

    assert [len(chunk) for chunk in report._get_issue_detail_chunks()] == [2]


def test_collapsing_can_be_turned_on():
    report = report_for(
        [SimpleNamespace(filename="ae.xpt", record_count=2)],
        args=SimpleNamespace(collapse_duplicate_issues=True),
    )
    report.get_detailed_data = lambda excel: [detail_row("CORE-1", "ae.xpt", row) for row in (1, 2)]

    assert [len(chunk) for chunk in report._get_issue_detail_chunks()] == [1]


def test_details_beyond_one_sheet_overflow_onto_numbered_sheets(monkeypatch):
    monkeypatch.setattr(ExcelReport, "_EXCEL_MAX_DATA_ROWS", 2)
    report = report_for([], args=SimpleNamespace(collapse_duplicate_issues=False))
    report.get_detailed_data = lambda excel: [detail_row(f"CORE-{row}", "ae.xpt", row) for row in range(1, 6)]
    wb = open_template()

    chunks = report._get_issue_detail_chunks()
    report._write_issue_details(wb, chunks, MagicMock())

    assert [len(chunk) for chunk in chunks] == [2, 2, 1]
    assert [wb[name]["A2"].value for name in ("Issue Details", "Issue Details 2", "Issue Details 3")] == [
        "CORE-1",
        "CORE-3",
        "CORE-5",
    ]
    assert wb["Issue Details"]["D2"].value == "AE"
